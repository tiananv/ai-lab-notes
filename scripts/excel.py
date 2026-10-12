"""Create a fictional workbook and check an agent's edits against its baseline."""

import argparse
import csv
import json
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from scripts.common import cli, write_csv, write_json
from scripts.llm import sha256


def create_workbook(inputs: Path, output: Path) -> None:
    """Build the small input workbook with formulas, never model output."""
    if output.exists():
        raise ValueError("Workbook already exists; choose a new --output")
    with inputs.open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    if len(rows) != 3:
        raise ValueError("This kit requires exactly three input orders")
    workbook = Workbook()
    orders = workbook.active
    orders.title = "Orders"
    orders.append(["Item", "Quantity", "Unit price", "Line total"])
    for number, row in enumerate(rows, 2):
        orders.append([row["item"], int(row["quantity"]), float(row["unit_price"]),
                       f"=B{number}*C{number}"])
    summary = workbook.create_sheet("Summary")
    for row in (["Measure", "Value"], ["Subtotal", "=SUM(Orders!D2:D4)"],
                ["Fictional handling charge", "=B2*0.1"], ["Total", "=B2+B3"]):
        summary.append(row)
    for sheet in workbook:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5BD0")
        sheet.freeze_panes = "A2"
        for column in ("A", "B", "C", "D", "E"):
            sheet.column_dimensions[column].width = 28
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()


def _cells(path: Path) -> tuple[dict, dict, list[str]]:
    try:
        workbook = load_workbook(path, data_only=False)
    except BadZipFile:
        raise ValueError(f"{path.name}: not a readable .xlsx workbook") from None
    values, formulas = {}, {}
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    key = f"{sheet.title}!{cell.coordinate}"
                    values[key] = cell.value
                    if cell.data_type == "f":
                        formulas[key] = cell.value
    names = workbook.sheetnames
    workbook.close()
    return values, formulas, names


def check_workbook(baseline: Path, candidate: Path) -> dict:
    """Check exact formula text, requested edits, and all other populated cells."""
    original, formulas, sheets = _cells(baseline)
    current, current_formulas, current_sheets = _cells(candidate)
    if not formulas:
        raise ValueError("Baseline contains no formulas; a pass would be meaningless")
    changed = [cell for cell, formula in formulas.items() if current_formulas.get(cell) != formula]
    added = sorted(current_formulas.keys() - formulas.keys())
    expected = {"Orders!B3": 4, "Orders!E1": "Status", "Orders!E2": "Ready",
                "Orders!E3": "Ready", "Orders!E4": "Ready"}
    expected_values = {**original, **expected}
    unexpected = sorted(cell for cell in current.keys() | expected_values.keys()
                        if current.get(cell) != expected_values.get(cell))
    edits_ok = all(current.get(cell) == value for cell, value in expected.items())
    formula_ok = not changed and not added
    return {"baseline_sha256": sha256(baseline.read_bytes()),
            "candidate_sha256": sha256(candidate.read_bytes()),
            "formula_count": len(formulas), "formulas_unchanged": formula_ok,
            "changed_or_missing_formulas": changed, "unexpected_formulas": added,
            "requested_edits_correct": edits_ok, "unexpected_cell_changes": unexpected,
            "sheets_unchanged": sheets == current_sheets,
            "passed": formula_ok and edits_ok and not unexpected and sheets == current_sheets,
            "recalculated_totals_checked": None, "formatting_checked": None}


def main() -> int:
    """Create an input workbook or score an actual candidate workbook."""
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    create = sub.add_parser("create")
    create.add_argument("--inputs", type=Path, default=Path("kits/excel-formulas/inputs/orders.csv"))
    create.add_argument("--output", type=Path, required=True)
    check = sub.add_parser("check")
    check.add_argument("--baseline", type=Path, required=True)
    check.add_argument("--candidate", type=Path, required=True)
    check.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "create":
        create_workbook(args.inputs, args.output)
        print(f"Created input workbook: {args.output}")
        return 0
    result = check_workbook(args.baseline, args.candidate)
    write_json(args.output.with_suffix(".json"), result)
    row = {key: result[key] for key in ("formula_count", "formulas_unchanged",
           "requested_edits_correct", "sheets_unchanged", "passed")}
    row.update(recalculated_totals_checked="", formatting_checked="", reviewer_notes="")
    write_csv(args.output, list(row), [row])
    print(json.dumps(result))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    cli(main)
