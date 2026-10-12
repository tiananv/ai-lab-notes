"""Show that formulas, task edits and unrelated cells are independently checked."""

from pathlib import Path

from openpyxl import load_workbook, Workbook
import pytest

from scripts.excel import create_workbook, check_workbook


@pytest.fixture
def books(tmp_path):
    """Create only synthetic workbook fixtures in the test cache."""
    baseline, candidate = tmp_path / "baseline.xlsx", tmp_path / "candidate.xlsx"
    create_workbook(Path("kits/excel-formulas/inputs/orders.csv"), baseline)
    workbook = load_workbook(baseline)
    workbook["Orders"]["B3"] = 4
    for cell, value in (("E1", "Status"), ("E2", "Ready"), ("E3", "Ready"), ("E4", "Ready")):
        workbook["Orders"][cell] = value
    workbook.save(candidate)
    workbook.close()
    return baseline, candidate


def test_requested_edits_pass_but_recalculation_is_unchecked(books):
    """Verify requested edits pass but recalculation is unchecked."""
    result = check_workbook(*books)
    assert result["passed"] is True and result["formula_count"] == 6
    assert result["recalculated_totals_checked"] is None


def test_unedited_copy_does_not_pass(books):
    """Verify unedited copy does not pass."""
    result = check_workbook(books[0], books[0])
    assert result["formulas_unchanged"] is True
    assert result["passed"] is False


@pytest.mark.parametrize(("sheet", "cell", "value"), [
    ("Orders", "D2", 80), ("Orders", "D2", "=B2*C3"), ("Summary", "B2", None),
    ("Orders", "A2", "Changed item"), ("Summary", "C4", "=1+1"),
])
def test_damage_is_detected(books, sheet, cell, value):
    """Verify damage is detected."""
    workbook = load_workbook(books[1])
    workbook[sheet][cell] = value
    workbook.save(books[1])
    workbook.close()
    assert check_workbook(*books)["passed"] is False


def test_missing_sheet_is_detected(books):
    """Verify missing sheet is detected."""
    workbook = load_workbook(books[1])
    del workbook["Summary"]
    workbook.save(books[1])
    workbook.close()
    result = check_workbook(*books)
    assert not result["sheets_unchanged"] and not result["formulas_unchanged"]


def test_no_formula_baseline_cannot_pass(tmp_path):
    """Verify no formula baseline cannot pass."""
    path = tmp_path / "blank.xlsx"
    workbook = Workbook()
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="no formulas"):
        check_workbook(path, path)


def test_cannot_overwrite_existing_baseline(books):
    """Verify cannot overwrite existing baseline."""
    with pytest.raises(ValueError, match="already exists"):
        create_workbook(Path("kits/excel-formulas/inputs/orders.csv"), books[0])
